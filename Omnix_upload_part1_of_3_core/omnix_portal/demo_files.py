"""
Example (demo) data files as Excel workbooks, for the Data preparation page and the Help pages.

Each workbook is built on request from the demo CSV files that ship with the platform, so what a
user downloads is exactly the data the platform loads as its demo. Sheet 1 ("Data") holds the table
in the layout Omnix expects: the platforms read the FIRST sheet of an uploaded workbook. Sheet 2
("Read me") explains every column.

An item with "file" instead of "csv" is a ready-made workbook that ships with the platform (e.g. the
Fluxomics Compound Discoverer templates, which carry their own Sample_List and Instructions sheets);
it is offered exactly as it is, under its own file name.
"""

import functools
import io
import os
import zipfile

import pandas as pd

from . import config, help_content

# platform -> [(section, [item, ...]), ...]; item: id, title, what it is, source CSV, optional row filter
DEMO = {
    "metabolomics": [
        ("Matrix", [
            {"id": "untargeted_metabolomics_matrix", "title": "Untargeted metabolomics: peak area matrix",
             "about": "300 metabolites × 30 injections (24 samples in 4 groups + 6 QCs).",
             "csv": "sample_peak_area_matrix_untargeted.csv"},
            {"id": "targeted_metabolomics_matrix", "title": "Targeted metabolomics: peak area matrix",
             "about": "59 metabolites + the internal standard ISTD_D4-Alanine × 30 injections.",
             "csv": "sample_peak_area_matrix_targeted.csv"},
        ]),
        ("Metadata", [
            {"id": "metabolomics_metadata", "title": "Sample metadata",
             "about": "Sample, Group, IsQC and Batch for the 30 injections; fits both matrices above.",
             "csv": "sample_metadata_untargeted.csv"},
        ]),
        ("Row annotation", [
            {"id": "untargeted_metabolomics_annotation", "title": "Untargeted metabolomics: row annotations",
             "about": "HMDB ID, Class and Method for each of the 300 metabolites.",
             "csv": "metabolite_row_annotations_untargeted.csv"},
            {"id": "targeted_metabolomics_annotation", "title": "Targeted metabolomics: row annotations",
             "about": "HMDB ID, Class, Method and ISTD (the internal standard, ISTD_D4-Alanine).",
             "csv": "metabolite_row_annotations_targeted.csv"},
        ]),
    ],
    "lipidomics": [
        ("Matrix", [
            {"id": "untargeted_lipidomics_matrix", "title": "Untargeted lipidomics: peak area matrix",
             "about": "90 lipids × 30 injections (24 samples in 4 groups + 6 QCs).",
             "csv": "sample_peak_area_matrix_multimethod_untargeted_lipidomics.csv"},
            {"id": "targeted_lipidomics_matrix", "title": "Targeted lipidomics: peak area matrix",
             "about": "39 lipids + the lipid internal standard ISTD_PC(15:0/18:1-d7) × 30 injections.",
             "csv": "sample_peak_area_matrix_multimethod_targeted_lipidomics.csv"},
        ]),
        ("Metadata", [
            {"id": "lipidomics_metadata", "title": "Sample metadata",
             "about": "Sample, Group, IsQC and Batch for the 30 injections; fits both matrices above.",
             "csv": "sample_metadata_untargeted.csv"},
        ]),
        ("Row annotation", [
            {"id": "untargeted_lipidomics_annotation", "title": "Untargeted lipidomics: row annotations",
             "about": "HMDB ID, Class and Method for the untargeted lipids.",
             "csv": "metabolite_row_annotations_multimethod.csv", "where": ("Method", "Untargeted Lipidomics")},
            {"id": "targeted_lipidomics_annotation", "title": "Targeted lipidomics: row annotations",
             "about": "HMDB ID, Class, Method and ISTD (the internal standard, ISTD_PC(15:0/18:1-d7)).",
             "csv": "metabolite_row_annotations_multimethod.csv", "where": ("Method", "Targeted Lipidomics")},
        ]),
    ],
    "proteomics": [
        ("Matrix", [
            {"id": "labelfree_proteomics_matrix", "title": "Label-Free proteomics: protein intensity matrix",
             "about": "243 proteins (Protein, Gene) × 30 samples (24 in 4 groups + 6 QCs).",
             "csv": "sample_protein_intensity_matrix_labelfree.csv"},
            {"id": "tmt_proteomics_matrix", "title": "TMT proteomics: protein intensity matrix",
             "about": "60 proteins × 31 channels, including the pooled reference channel Reference_Pool.",
             "csv": "sample_protein_intensity_matrix_tmt.csv"},
        ]),
        ("Metadata", [
            {"id": "labelfree_proteomics_metadata", "title": "Label-Free proteomics: sample metadata",
             "about": "Sample, Group, IsQC and Batch for the 30 samples.",
             "csv": "sample_metadata_labelfree.csv"},
            {"id": "tmt_proteomics_metadata", "title": "TMT proteomics: sample metadata",
             "about": "As Label-Free, plus the Reference_Pool channel (Group = Reference).",
             "csv": "sample_metadata_tmt.csv"},
        ]),
        ("Row annotation", [
            {"id": "proteomics_annotation", "title": "Protein row annotations",
             "about": "Protein, Gene, Method and Pathway for 200 proteins (Label-Free and TMT).",
             "csv": "protein_row_annotations_richdemo.csv"},
        ]),
    ],
    "fluxomics": [
        ("Isotopologue table", [
            {"id": "fluxomics_isotopologue_table_synthetic", "title": "Isotopologue table: synthetic example",
             "about": "63 metabolites, 373 isotopologue rows (M+0 … M+n) × 20 samples: shControl vs GeneKD, "
                      "[U-13C6]-glucose, 2 / 8 / 24 h, 3 replicates, one unlabeled reference and one blank. "
                      "Simulated data; the demo the platform loads.",
             "csv": "example_data/generic/synthetic_isotopologue_table.csv"},
            {"id": "fluxomics_isotopologue_table_with_12c_reference", "title": "Template: table with 12C reference",
             "about": "Generic layout with the unlabeled (12C) reference column: 18 labeled samples + Unlabeled_Ref_1.",
             "csv": "example_data/templates/generic/Generic_Isotopologue_Table_with_12C_Ref.csv"},
            {"id": "fluxomics_isotopologue_table_without_12c_reference",
             "title": "Template: table without 12C reference",
             "about": "The same table without a reference column (correction then uses theory only).",
             "csv": "example_data/templates/generic/Generic_Isotopologue_Table_without_12C_Ref.csv"},
        ]),
        ("Sample metadata", [
            {"id": "fluxomics_sample_metadata_synthetic", "title": "Sample metadata: synthetic example",
             "about": "Sample, Type, Group, Treatment, Time, TimeUnit, Replicate and Include for the 20 samples.",
             "csv": "example_data/generic/synthetic_sample_metadata.csv"},
            {"id": "fluxomics_sample_metadata_template", "title": "Template: sample metadata",
             "about": "Fits both template tables (a row without a matching column is ignored).",
             "csv": "example_data/templates/generic/Generic_Sample_Metadata.csv"},
        ]),
        ("Compound Discoverer export", [
            {"id": "fluxomics_cd_sil_rp_pos", "title": "Compound Discoverer SIL export: RP_Pos template",
             "about": "SIL export on the first sheet, plus Sample_List and Instructions sheets. Upload it together "
                      "with the HILIC_Neg template, or replace its rows with your own export.",
             "file": "example_data/templates/cd_sil/CD_SIL_Template_RP_Pos.xlsx"},
            {"id": "fluxomics_cd_sil_hilic_neg", "title": "Compound Discoverer SIL export: HILIC_Neg template",
             "about": "The second method of the same experiment (shControl vs GeneKD, 12 / 18 / 40 h).",
             "file": "example_data/templates/cd_sil/CD_SIL_Template_HILIC_Neg.xlsx"},
        ]),
    ],
    "transcriptomics": [
        ("Matrix", [
            {"id": "transcriptomics_counts_matrix", "title": "Bulk RNA-seq: raw counts matrix",
             "about": "2,178 genes × 46 samples (4 diagnosis groups of 10 + 6 QCs).",
             "csv": "sample_counts_matrix_richdemo.csv"},
        ]),
        ("Metadata", [
            {"id": "transcriptomics_metadata", "title": "Clinical metadata",
             "about": "Sample, Group, IsQC, Batch plus Diagnosis, Age, Gender, Treatment, Ethnicity, Body weight.",
             "csv": "sample_metadata_richdemo.csv"},
        ]),
        ("Row annotation", [
            {"id": "transcriptomics_annotation", "title": "Gene row annotations",
             "about": "Gene and Pathway for all 2,178 genes.",
             "csv": "gene_row_annotations_richdemo.csv"},
        ]),
    ],
}

_SECTION_SPEC = {"Matrix": 0, "Metadata": 1, "Row annotation": 2,    # index into help_content.FILES
                 "Isotopologue table": 0, "Sample metadata": 1, "Compound Discoverer export": 2}

# Data preparation page: the line above a platform's files (default: the two-sheet workbook layout)
INTRO = {
    "fluxomics": "Generic files come as workbooks with two sheets: **Data**, the table to upload (Omnix reads the "
                 "first sheet), and **Read me**, which explains each column. The Compound Discoverer templates are "
                 "the original export workbooks, with their own *Sample_List* and *Instructions* sheets.",
}


def filename(platform, item):
    """e.g. Omnix_Untargeted_Lipidomics_Matrix.xlsx (a ready-made workbook keeps its own name)"""
    if item.get("file"):
        return os.path.basename(item["file"])
    fix = {"Labelfree": "LabelFree", "Tmt": "TMT", "12C": "12C", "Cd": "CD", "Sil": "SIL", "Rp": "RP",
           "Hilic": "HILIC"}
    return "Omnix_" + "_".join(fix.get(w.title(), w.title()) for w in item["id"].split("_")) + ".xlsx"


def items(platform):
    return [(section, it) for section, its in DEMO[platform] for it in its]


def _table(app_dir, item):
    df = pd.read_csv(os.path.join(app_dir, item["csv"]))
    if item.get("where"):
        col, val = item["where"]
        df = df[df[col] == val].reset_index(drop=True)
    return df.dropna(axis=1, how="all")          # e.g. the ISTD column, which untargeted files leave empty


@functools.lru_cache(maxsize=64)
def _workbook(platform, item_id, app_dir, mtime):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    section, item = next((s, it) for s, it in items(platform) if it["id"] == item_id)
    if item.get("file"):                         # ready-made workbook: offered as it is
        path = os.path.join(app_dir, item["file"])
        with open(path, "rb") as fh:
            return fh.read(), pd.read_excel(path, sheet_name=0)
    df = _table(app_dir, item)
    spec = help_content.FILES[platform][_SECTION_SPEC[section]]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Data", index=False)
        ws = xw.sheets["Data"]
        head_fill, head_font = PatternFill("solid", fgColor="00695C"), Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "B2"
        for i, col in enumerate(df.columns, 1):
            width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 42)

        rows = [[f"{config.BRAND}™ example file: {item['title']}"], [item["about"]], [],
                ["This workbook is the demo dataset the platform itself uses. Keep your data on the FIRST "
                 "sheet (\"Data\"): Omnix reads the first sheet of an uploaded workbook."],
                ["You can replace the values with your own and upload the file as it is (.xlsx), or save it as CSV."],
                [], ["Column", "Required", "What goes here", "Example"]]
        rows += [list(r) for r in spec["columns"]]
        rows += [[], ["Before you upload"]] + [[c.replace("**", "")] for c in help_content.checklist(platform)]
        pd.DataFrame(rows).to_excel(xw, sheet_name="Read me", index=False, header=False)
        rm = xw.sheets["Read me"]
        rm["A1"].font = Font(bold=True, size=14, color="00574D")
        hdr = 7
        for cell in rm[hdr]:
            cell.fill, cell.font = head_fill, head_font
        rm.cell(row=hdr + len(spec["columns"]) + 2, column=1).font = Font(bold=True, color="00574D")
        for letter, w in zip("ABCD", (34, 14, 80, 26)):
            rm.column_dimensions[letter].width = w
        for row in rm.iter_rows(min_row=hdr + 1, max_row=hdr + len(spec["columns"])):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    return buf.getvalue(), df


def workbook(platform, item_id, app_dir):
    """(xlsx bytes, DataFrame) for one example file."""
    it = next(it for _, it in items(platform) if it["id"] == item_id)
    path = os.path.join(app_dir, it.get("file") or it["csv"])
    return _workbook(platform, item_id, app_dir, os.path.getmtime(path))


def zip_all(platform, app_dir):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for _, it in items(platform):
            data, _ = workbook(platform, it["id"], app_dir)
            zf.writestr(filename(platform, it), data)
    return buf.getvalue()
